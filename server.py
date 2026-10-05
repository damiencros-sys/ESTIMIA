from datetime import datetime

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from pathlib import Path
import os, json, uuid, urllib.request, urllib.error, urllib.parse, mimetypes, re, math, io

BASE=Path(__file__).resolve().parent
DATA=BASE/"data"; DATA.mkdir(exist_ok=True)
app=FastAPI(title="ESTIM'IA V2")
app.mount("/static",StaticFiles(directory=BASE/"static"),name="static")

@app.get("/")
def home(): return FileResponse(BASE/"static"/"index.html")


def get_json(url, timeout=20):
    req=urllib.request.Request(url,headers={"User-Agent":"ESTIMIA/10"})
    with urllib.request.urlopen(req,timeout=timeout) as r:
        return json.loads(r.read())

@app.get("/api/geocode")
def geocode(q:str, postcode:str="", city:str="", citycode:str=""):
    """Vérification stricte dans le référentiel national d'adresse alimenté par les BAL.
    Une réponse située dans une autre commune n'est jamais renvoyée.
    """
    q=q.strip(); postcode=postcode.strip(); city=city.strip(); citycode=citycode.strip()
    if len(q)<3: return {"results":[]}
    params={"q":q,"limit":15}
    if postcode: params["postcode"]=postcode
    if citycode: params["citycode"]=citycode
    url="https://data.geopf.fr/geocodage/search/?"+urllib.parse.urlencode(params)
    try:
        data=get_json(url); out=[]; wanted_city=_norm_name(city)
        for f in data.get("features",[]):
            pr=f.get("properties",{}); pc=str(pr.get("postcode") or "")
            cc=str(pr.get("citycode") or pr.get("city_code") or "")
            cn=pr.get("city") or pr.get("municipality") or ""
            if isinstance(cn,list): cn=cn[0] if cn else ""
            if isinstance(pc,list): pc=pc[0] if pc else ""
            if isinstance(cc,list): cc=cc[0] if cc else ""
            # Verrou commune : aucune proposition hors commune/CP demandé.
            if citycode and cc != citycode: continue
            if postcode and pc != postcode: continue
            if wanted_city and _norm_name(str(cn)) != wanted_city: continue
            out.append({
                "label":pr.get("label") or pr.get("name") or "", "postcode":pc,
                "city":cn, "citycode":cc, "score":pr.get("score"),
                "parcels":pr.get("cad_parcelles") or pr.get("parcelles") or [],
                "coordinates":(f.get("geometry") or {}).get("coordinates")
            })
        return {"results":out}
    except Exception:
        raise HTTPException(502,"Le service public d'adresses ne répond pas pour le moment.")


def _norm_name(s):
    import unicodedata
    s=unicodedata.normalize("NFD",s or "")
    s="".join(c for c in s if unicodedata.category(c)!="Mn").lower()
    return re.sub(r"[^a-z0-9]+"," ",s).strip()

def _resolve_commune(q):
    q=(q or "").strip()
    if re.fullmatch(r"\d{5}",q): return {"name":q,"code":q,"postcode":""}
    mcp=re.search(r"\b(\d{5})\b",q); cp=mcp.group(1) if mcp else ""
    wanted=re.sub(r"\b\d{5}\b"," ",q).strip(" ,.-")
    if cp:
        try:
            rows=get_json(f"https://apicarto.ign.fr/api/codes-postaux/communes/{cp}")
            if isinstance(rows,dict): rows=rows.get("communes") or rows.get("results") or []
            nw=_norm_name(wanted); choices=[]
            for row in rows or []:
                name=row.get("nomCommune") or row.get("nom") or row.get("libelle") or row.get("name") or ""
                code=row.get("codeCommune") or row.get("code_insee") or row.get("insee") or row.get("code") or ""
                if name and code:
                    score=4 if _norm_name(name)==nw else (2 if nw and (nw in _norm_name(name) or _norm_name(name) in nw) else 0)
                    choices.append((score,{"name":name,"code":str(code),"postcode":cp}))
            if choices:
                choices.sort(key=lambda x:x[0],reverse=True)
                if choices[0][0]>0 or len(choices)==1:return choices[0][1]
        except Exception: pass
    for query in [q,wanted]:
        if not query: continue
        try:
            u="https://data.geopf.fr/geocodage/search/?"+urllib.parse.urlencode({"q":query,"limit":10})
            fs=get_json(u).get("features",[]); nw=_norm_name(wanted or query); choices=[]
            for f in fs:
                pr=f.get("properties",{})
                code=pr.get("citycode") or pr.get("city_code") or ""
                name=pr.get("city") or pr.get("municipality") or pr.get("name") or ""
                pc=pr.get("postcode") or cp or ""
                if isinstance(name,list): name=name[0] if name else ""
                if isinstance(code,list): code=code[0] if code else ""
                if isinstance(pc,list): pc=pc[0] if pc else ""
                if not name or not code: continue
                score=4 if _norm_name(name)==nw else (2 if nw and (nw in _norm_name(name) or _norm_name(name) in nw) else 0)
                if cp and str(pc)==cp: score+=3
                choices.append((score,{"name":name,"code":str(code),"postcode":str(pc)}))
            if choices:
                choices.sort(key=lambda x:x[0],reverse=True); return choices[0][1]
        except Exception: pass
    return None

@app.get("/api/commune")
def commune_lookup(q:str):
    return {"result":_resolve_commune(q)}


def _section_candidates(section:str):
    s=re.sub(r"[^A-Za-z0-9]","",(section or "").upper().strip())
    if not s:return []
    vals=[]
    if len(s)==1 and s.isalpha(): vals.append("0"+s)
    if s.isdigit(): vals.append(s.zfill(2))
    vals.append(s)
    return list(dict.fromkeys(vals))

def _fetch_parcel(code_insee:str, section:str, numero:str):
    no=re.sub(r"\D","",str(numero or "")).zfill(4)
    if not no.strip("0"): raise HTTPException(400,"Numéro de parcelle invalide.")
    had_response=False; last_error=None
    for sec in _section_candidates(section):
        for source in ("PCI",None):
            params={"code_insee":code_insee,"section":sec,"numero":no}
            if source:params["source_ign"]=source
            try:
                d=get_json("https://apicarto.ign.fr/api/cadastre/parcelle?"+urllib.parse.urlencode(params));had_response=True
                fs=d.get("features",[])
                if fs:return fs[0],sec,no
            except Exception as e:last_error=e
    if not had_response and last_error:raise last_error
    return None,(_section_candidates(section) or [section])[0],no

@app.get("/api/cadastre")
def cadastre(commune:str, section:str, numero:str):
    commune=commune.strip()
    resolved={"name":commune,"code":commune,"postcode":""} if re.fullmatch(r"\d{5}",commune) else _resolve_commune(commune)
    if not resolved or not resolved.get("code"):raise HTTPException(400,"Commune non reconnue.")
    code_insee=str(resolved["code"]);commune_name=resolved.get("name") or commune
    try:
        f,sec_api,no=_fetch_parcel(code_insee,section,numero);display_sec=re.sub(r"^0(?=[A-Z]$)","",sec_api)
        if not f:return {"found":False,"commune":commune_name,"code_insee":code_insee,"section":display_sec,"numero":no}
        return {"found":True,"commune":commune_name,"code_insee":code_insee,"section":display_sec,"numero":no,"properties":f.get("properties",{}),"geometry":f.get("geometry")}
    except HTTPException:raise
    except Exception:raise HTTPException(502,"Le service cadastral officiel est momentanément indisponible. Réessaie dans quelques secondes.")

def _parcel_feature(commune:str, section:str, numero:str):
    commune=(commune or "").strip()
    resolved={"name":commune,"code":commune,"postcode":""} if re.fullmatch(r"\d{5}",commune) else _resolve_commune(commune)
    if not resolved or not resolved.get("code"):raise HTTPException(400,"Commune non reconnue.")
    code_insee=str(resolved["code"]);commune_name=resolved.get("name") or commune
    f,sec_api,no=_fetch_parcel(code_insee,section,numero);display_sec=re.sub(r"^0(?=[A-Z]$)","",sec_api)
    return f,commune_name,code_insee,display_sec,no

def _geom_rings(g):
    if not g:return []
    c=g.get("coordinates") or []
    if g.get("type")=="Polygon": return c
    if g.get("type")=="MultiPolygon":
        out=[]
        for poly in c: out.extend(poly)
        return out
    return []

@app.get("/api/cadastre/plan")
def cadastre_plan(commune:str, section:str, numero:str):
    try:
        f,name,code,sec,no=_parcel_feature(commune,section,numero)
        if not f: raise HTTPException(404,"Parcelle non trouvée.")
        rings=_geom_rings(f.get("geometry")); pts=[p for ring in rings for p in ring if isinstance(p,list) and len(p)>=2]
        if not pts: raise HTTPException(404,"Géométrie cadastrale indisponible.")
        xs=[float(p[0]) for p in pts]; ys=[float(p[1]) for p in pts]
        xmin,xmax=min(xs),max(xs); ymin,ymax=min(ys),max(ys); dx=max(xmax-xmin,1e-9); dy=max(ymax-ymin,1e-9)
        W,H,P=900,520,45; scale=min((W-2*P)/dx,(H-2*P)/dy)
        def xy(p): return f"{P+(float(p[0])-xmin)*scale:.1f},{H-P-(float(p[1])-ymin)*scale:.1f}"
        paths=["M "+" L ".join(xy(p) for p in ring)+" Z" for ring in rings if len(ring)>=3]
        label=f"{name} - Section {sec} - Parcelle {int(no)}"; path_data=" ".join(paths)
        svg=("<svg xmlns='http://www.w3.org/2000/svg' width='900' height='520' viewBox='0 0 900 520'>"
             "<rect width='100%' height='100%' fill='white'/>"
             f"<text x='45' y='28' font-family='Arial,sans-serif' font-size='18' font-weight='700'>{label}</text>"
             f"<path d='{path_data}' fill='#f4f4f4' stroke='#222' stroke-width='3'/>"
             "<text x='45' y='505' font-family='Arial,sans-serif' font-size='12'>Contour cadastral officiel IGN - representation indicative</text></svg>")
        return Response(content=svg,media_type="image/svg+xml",headers={"Cache-Control":"no-store"})
    except HTTPException: raise
    except Exception: raise HTTPException(502,"Impossible de générer le plan cadastral pour le moment.")

def _lonlat_points(geometry):
    rings=_geom_rings(geometry)
    return [(float(p[0]),float(p[1])) for ring in rings for p in ring if isinstance(p,list) and len(p)>=2]

def _tile_xy(lon,lat,z):
    lat=max(min(lat,85.05112878),-85.05112878)
    n=2**z
    x=(lon+180.0)/360.0*n
    y=(1.0-math.asinh(math.tan(math.radians(lat)))/math.pi)/2.0*n
    return x,y

def _mercator(lon,lat):
    lat=max(min(float(lat),85.05112878),-85.05112878); lon=float(lon)
    R=6378137.0
    return R*math.radians(lon), R*math.log(math.tan(math.pi/4+math.radians(lat)/2))

def _context_map_png(commune,section,numero,width=760,height=520):
    from PIL import Image, ImageDraw
    f,name,code,sec,no=_parcel_feature(commune,section,numero)
    if not f: raise HTTPException(404,"Parcelle non trouvée.")
    geom=f.get("geometry"); rings=_geom_rings(geom)
    mpts=[_mercator(p[0],p[1]) for ring in rings for p in ring if len(p)>=2]
    if not mpts: raise HTTPException(404,"Géométrie cadastrale indisponible.")
    xs=[p[0] for p in mpts]; ys=[p[1] for p in mpts]
    cx=(min(xs)+max(xs))/2; cy=(min(ys)+max(ys))/2
    span=max(max(xs)-min(xs),max(ys)-min(ys),55.0)*5.0
    # Emprise contextuelle, avec correction du ratio de l'image.
    aspect=width/height
    half_y=span/2; half_x=half_y*aspect
    bbox=(cx-half_x,cy-half_y,cx+half_x,cy+half_y)
    layers="AMORCES_CAD,LIEUDIT,CP.CadastralParcel,SUBFISCAL,CLOTURE,DETAIL_TOPO,HYDRO,VOIE_COMMUNICATION,BU.Building,BORNE_REPERE"
    params={"service":"WMS","version":"1.3","request":"GetMap","layers":layers,"styles":"",
            "format":"image/png","crs":"EPSG:3857","bbox":",".join(f"{v:.3f}" for v in bbox),
            "width":str(width),"height":str(height),"language":"fre"}
    url=f"https://inspire.cadastre.gouv.fr/scpc/{code}.wms?"+urllib.parse.urlencode(params)
    try:
        req=urllib.request.Request(url,headers={"User-Agent":"ESTIMIA/13.8"})
        with urllib.request.urlopen(req,timeout=20) as r:
            raw=r.read(); ctype=r.headers.get("Content-Type","")
        if "image" not in ctype.lower(): raise ValueError("Réponse WMS non image")
        out=Image.open(io.BytesIO(raw)).convert("RGB")
    except Exception:
        raise HTTPException(502,"Le plan cadastral officiel DGFiP est momentanément indisponible.")
    draw=ImageDraw.Draw(out,"RGBA")
    xmin,ymin,xmax,ymax=bbox
    def px(pt):
        x,y=_mercator(pt[0],pt[1])
        return ((x-xmin)/(xmax-xmin)*width, height-(y-ymin)/(ymax-ymin)*height)
    for ring in rings:
        poly=[px(p) for p in ring if len(p)>=2]
        if len(poly)>=3: draw.polygon(poly,fill=(255,35,35,55),outline=(235,0,0,255),width=4)
    mx=(cx-xmin)/(xmax-xmin)*width; my=height-(cy-ymin)/(ymax-ymin)*height
    draw.ellipse((mx-10,my-10,mx+10,my+10),fill=(35,190,70,255),outline=(255,255,255,255),width=3)
    draw.rectangle((0,height-28,width,height),fill=(255,255,255,225))
    draw.text((8,height-22),f"{name} - Section {sec} - Parcelle {int(no)} | Plan cadastral officiel DGFiP",fill=(20,20,20,255))
    bio=io.BytesIO(); out.save(bio,format="PNG",optimize=True); bio.seek(0)
    return bio,name,code,sec,no


def _multi_context_map_png(commune, refs, width=760, height=520):
    from PIL import Image, ImageDraw
    refs=[x for x in refs if x.get("section") and x.get("numero")]
    if not refs: raise HTTPException(400,"Aucune parcelle fournie.")
    features=[]; commune_name=""; code_insee=""
    for ref in refs:
        f,name,code,sec,no=_parcel_feature(commune,ref["section"],ref["numero"])
        if not f: raise HTTPException(404,f"Parcelle {sec} {int(no)} non trouvée.")
        commune_name=commune_name or name; code_insee=code_insee or code
        features.append((f,sec,no))
    all_m=[]
    for f,sec,no in features:
        for ring in _geom_rings(f.get("geometry")):
            all_m.extend(_mercator(p[0],p[1]) for p in ring if len(p)>=2)
    if not all_m: raise HTTPException(404,"Géométrie cadastrale indisponible.")
    xs=[p[0] for p in all_m]; ys=[p[1] for p in all_m]
    cx=(min(xs)+max(xs))/2; cy=(min(ys)+max(ys))/2
    span=max(max(xs)-min(xs),max(ys)-min(ys),55.0)*1.55
    aspect=width/height; half_y=span/2; half_x=max(span*aspect/2,(max(xs)-min(xs))*0.65)
    half_y=max(half_y,(max(ys)-min(ys))*0.65)
    bbox=(cx-half_x,cy-half_y,cx+half_x,cy+half_y)
    layers="AMORCES_CAD,LIEUDIT,CP.CadastralParcel,SUBFISCAL,CLOTURE,DETAIL_TOPO,HYDRO,VOIE_COMMUNICATION,BU.Building,BORNE_REPERE"
    params={"service":"WMS","version":"1.3","request":"GetMap","layers":layers,"styles":"",
            "format":"image/png","crs":"EPSG:3857","bbox":",".join(f"{v:.3f}" for v in bbox),
            "width":str(width),"height":str(height),"language":"fre"}
    url=f"https://inspire.cadastre.gouv.fr/scpc/{code_insee}.wms?"+urllib.parse.urlencode(params)
    try:
        req=urllib.request.Request(url,headers={"User-Agent":"ESTIMIA/13.34"})
        with urllib.request.urlopen(req,timeout=20) as r:
            raw=r.read(); ctype=r.headers.get("Content-Type","")
        if "image" not in ctype.lower(): raise ValueError("Réponse WMS non image")
        out=Image.open(io.BytesIO(raw)).convert("RGB")
    except Exception:
        raise HTTPException(502,"Le plan cadastral officiel DGFiP est momentanément indisponible.")
    draw=ImageDraw.Draw(out,"RGBA"); xmin,ymin,xmax,ymax=bbox
    def px(pt):
        x,y=_mercator(pt[0],pt[1])
        return ((x-xmin)/(xmax-xmin)*width, height-(y-ymin)/(ymax-ymin)*height)
    for idx,(f,sec,no) in enumerate(features,1):
        centers=[]
        for ring in _geom_rings(f.get("geometry")):
            poly=[px(p) for p in ring if len(p)>=2]
            if len(poly)>=3:
                draw.polygon(poly,fill=(255,35,35,62),outline=(220,0,0,255),width=4)
                centers.extend(poly)
        if centers:
            mx=sum(p[0] for p in centers)/len(centers); my=sum(p[1] for p in centers)/len(centers)
            label=f"{sec} {int(no)}"
            draw.ellipse((mx-11,my-11,mx+11,my+11),fill=(35,170,70,255),outline=(255,255,255,255),width=2)
            draw.text((mx+14,my-8),label,fill=(15,60,25,255),stroke_width=2,stroke_fill=(255,255,255,240))
    refs_label=" • ".join(f"{sec} {int(no)}" for _,sec,no in features)
    draw.rectangle((0,height-30,width,height),fill=(255,255,255,230))
    draw.text((8,height-23),f"{commune_name} — {refs_label} | Plan cadastral officiel DGFiP",fill=(20,20,20,255))
    bio=io.BytesIO(); out.save(bio,format="PNG",optimize=True); bio.seek(0)
    return bio,commune_name,code_insee

@app.get("/api/cadastre/multi-map.png")
def cadastre_multi_map(commune:str, refs:str):
    parsed=[]
    for token in refs.split(","):
        token=token.strip()
        m=re.fullmatch(r"([A-Za-z0-9]{1,3})\s*:\s*(\d{1,4})",token)
        if m: parsed.append({"section":m.group(1).upper(),"numero":m.group(2)})
    bio,_,_=_multi_context_map_png(commune,parsed)
    return StreamingResponse(bio,media_type="image/png",headers={"Cache-Control":"no-store"})

@app.get("/api/cadastre/map.png")
def cadastre_context_map(commune:str, section:str, numero:str):
    bio,_,_,_,_=_context_map_png(commune,section,numero)
    return StreamingResponse(bio,media_type="image/png",headers={"Cache-Control":"no-store"})

class WordPayload(BaseModel):
    owner:str=""
    phone:str=""
    address:str=""
    commune:str=""
    section:str=""
    parcel:str=""
    property_type:str=""
    surface:str=""
    surface_carrez:str=""
    sections:list=[]
    parcels:list=[]
    facts:str=""
    photos:list=[]
    generated_date:str=""
    theme_color:str=""

def _property_color(t):
    s=(t or "").lower()
    if any(x in s for x in ["appartement","studio"]): return "2563A6"
    if "terrain" in s: return "B7791F"
    if "immeuble" in s: return "8B2F45"
    if any(x in s for x in ["local","commerce","commercial","professionnel","bureau"]): return "7651A8"
    if any(x in s for x in ["garage","parking","box"]): return "64748B"
    if any(x in s for x in ["maison","villa","pavillon"]): return "2F7D5A"
    return "49697D"

@app.post("/api/word")
def word_export(p:WordPayload):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.45); sec.bottom_margin=Inches(.45); sec.left_margin=Inches(.55); sec.right_margin=Inches(.55)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5)
    styles['Title'].font.name='Aptos Display'; styles['Title'].font.size=Pt(20); styles['Title'].font.bold=True

    def shade(cell,fill):
        tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
    def margins(cell,top=70,start=90,bottom=70,end=90):
        tc=cell._tc.get_or_add_tcPr(); mar=tc.first_child_found_in('w:tcMar')
        if mar is None: mar=OxmlElement('w:tcMar'); tc.append(mar)
        for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
            el=OxmlElement('w:'+tag); el.set(qn('w:w'),str(val)); el.set(qn('w:type'),'dxa'); mar.append(el)
    def keep_table(table):
        for row in table.rows:
            trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)
    def label_value_table(rows,widths=(2.15,4.75)):
        table=doc.add_table(rows=0,cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
        for k,v in rows:
            if not v: continue
            cells=table.add_row().cells; cells[0].text=str(k); cells[1].text=str(v)
            cells[0].width=Inches(widths[0]); cells[1].width=Inches(widths[1])
            cells[0].vertical_alignment=cells[1].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cells[0],'EAF3F1')
            for c in cells: margins(c)
            for r in cells[0].paragraphs[0].runs: r.bold=True; r.font.size=Pt(9)
            for r in cells[1].paragraphs[0].runs: r.font.size=Pt(9)
        keep_table(table); return table
    property_color=_property_color(p.property_type)
    def heading(text):
        para=doc.add_paragraph(); para.paragraph_format.space_before=Pt(7); para.paragraph_format.space_after=Pt(3); para.paragraph_format.keep_with_next=True
        r=para.add_run(text); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(property_color)
        return para

    title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_after=Pt(4)
    r=title.add_run('FICHE DE VISITE IMMOBILIÈRE'); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor.from_string(property_color)
    sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after=Pt(7)
    rr=sub.add_run('ESTIM’IA — fiche de relevé'); rr.italic=True; rr.font.size=Pt(9); rr.font.color.rgb=RGBColor(0x66,0x66,0x66)
    datep=doc.add_paragraph(); datep.alignment=WD_ALIGN_PARAGRAPH.CENTER; datep.paragraph_format.space_after=Pt(7)
    dr=datep.add_run('Document généré le '+(p.generated_date or datetime.now().strftime('%d/%m/%Y'))); dr.font.size=Pt(8.5); dr.font.color.rgb=RGBColor(0x66,0x66,0x66)

    heading('DOSSIER')
    meta=[('Propriétaire',p.owner),('Téléphone',p.phone),('Adresse',p.address),('Commune',p.commune),
          ('Cadastre',(' • '.join([str(x.get('section','')).upper()+' '+str(int(re.sub(r'\\D','',str(x.get('numero',''))) or '0')) for x in p.parcels if x.get('section') and x.get('numero')]) if p.parcels else (('Section '+p.section.upper()+' — Parcelle '+str(int(re.sub(r'\\D','',p.parcel) or '0'))) if p.section and p.parcel else ''))),
          ('Type',p.property_type),('Surface habitable',(p.surface+' m²') if p.surface else ''),('Surface Carrez',(p.surface_carrez+' m²') if p.surface_carrez else '')]
    label_value_table(meta)

    refs=[x for x in p.parcels if isinstance(x,dict) and x.get('section') and x.get('numero')]
    if not refs and p.section and p.parcel: refs=[{"section":p.section,"numero":p.parcel}]
    if p.commune and refs:
        try:
            if len(refs)>1:
                img,_,_=_multi_context_map_png(p.commune,refs,760,500)
            else:
                img,_,_,_,_=_context_map_png(p.commune,refs[0]["section"],refs[0]["numero"],760,500)
            heading('PLAN CADASTRAL')
            doc.add_picture(img,width=Inches(5.35)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
            refs_txt=' • '.join(str(x["section"]).upper()+' '+str(int(re.sub(r"\\D","",str(x["numero"])) or "0")) for x in refs)
            cap=doc.add_paragraph(f'{p.commune} — Parcelles : {refs_txt}')
            cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after=Pt(4)
            for r in cap.runs: r.font.size=Pt(8); r.font.italic=True
        except Exception as e:
            note=doc.add_paragraph('Plan cadastral temporairement indisponible lors de la génération du document.')
            note.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in note.runs: r.font.size=Pt(8); r.font.italic=True

    if p.sections:
        heading('RELEVÉ DE VISITE')
        for secdata in p.sections:
            title=str(secdata.get('title','')).strip()
            rows=secdata.get('rows') or []
            rows=[[str(x[0]),str(x[1])] for x in rows if isinstance(x,(list,tuple)) and len(x)>=2 and str(x[1]).strip()]
            if not rows: continue
            heading(title)
            label_value_table(rows)
    elif p.facts.strip():
        heading('RELEVÉ DE VISITE')
        for line in [x.strip() for x in p.facts.splitlines() if x.strip()]:
            doc.add_paragraph(line)


    # Photographies ajoutées au dossier. Les contrôles de l'interface ne sont jamais exportés.
    if p.photos:
        import base64
        heading('PHOTOGRAPHIES DU BIEN')
        valid=[]
        for item in p.photos[:30]:
            try:
                data=str(item.get('data',''))
                if ',' not in data: continue
                raw=base64.b64decode(data.split(',',1)[1])
                if len(raw)>12_000_000: continue
                valid.append((item.get('name','Photo'),raw))
            except Exception:
                continue
        for i in range(0,len(valid),2):
            table=doc.add_table(rows=1,cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER
            for j in range(2):
                cell=table.rows[0].cells[j]
                if i+j>=len(valid): continue
                name,raw=valid[i+j]
                try:
                    bio=io.BytesIO(raw)
                    pcell=cell.paragraphs[0]; pcell.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    run=pcell.add_run(); run.add_picture(bio,width=Inches(3.0))
                    cap=cell.add_paragraph(str(name)); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for rr in cap.runs: rr.font.size=Pt(7); rr.font.italic=True
                except Exception:
                    pass

    # footer + page number field
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fr=footer.add_run('ESTIM’IA  •  Fiche de visite  •  Page '); fr.font.size=Pt(8); fr.font.color.rgb=RGBColor(0x77,0x77,0x77)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

    out=io.BytesIO(); doc.save(out); out.seek(0)
    safe=re.sub(r'[^A-Za-z0-9_-]+','_',p.owner or p.address or 'bien').strip('_') or 'bien'
    return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':f'attachment; filename="Fiche_visite_{safe}.docx"'})


@app.post("/api/pdf")
def pdf_export(p:WordPayload):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, KeepTogether, PageBreak
    import base64
    theme=colors.HexColor('#'+_property_color(p.property_type))
    out=io.BytesIO(); doc=SimpleDocTemplate(out,pagesize=A4,rightMargin=14*mm,leftMargin=14*mm,topMargin=13*mm,bottomMargin=14*mm)
    styles=getSampleStyleSheet()
    title=ParagraphStyle('TitleX',parent=styles['Title'],fontName='Helvetica-Bold',fontSize=17,leading=20,textColor=theme,alignment=TA_CENTER,spaceAfter=3*mm)
    date_style=ParagraphStyle('DateX',parent=styles['Normal'],fontSize=8,textColor=colors.HexColor('#666666'),alignment=TA_CENTER,spaceAfter=4*mm)
    h=ParagraphStyle('HX',parent=styles['Heading2'],fontName='Helvetica-Bold',fontSize=10.5,leading=13,textColor=theme,spaceBefore=4*mm,spaceAfter=2*mm)
    normal=ParagraphStyle('NX',parent=styles['Normal'],fontSize=8.5,leading=11)
    cap=ParagraphStyle('CapX',parent=normal,fontSize=7.5,textColor=colors.HexColor('#555555'),alignment=TA_CENTER)
    story=[Paragraph('FICHE DE VISITE IMMOBILIERE',title),Paragraph('ESTIM’IA - fiche de releve',date_style),Paragraph('Document genere le '+(p.generated_date or datetime.now().strftime('%d/%m/%Y')),date_style)]
    def add_table(rows):
        data=[]
        for k,v in rows:
            if str(v).strip(): data.append([Paragraph(str(k),normal),Paragraph(str(v),normal)])
        if not data:return
        t=Table(data,colWidths=[52*mm,112*mm],repeatRows=0,hAlign='CENTER')
        t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.HexColor('#CCD5DB')),('BACKGROUND',(0,0),(0,-1),colors.HexColor('#F2F6F5')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),('TEXTCOLOR',(0,0),(0,-1),theme)]));story.append(t)
    story.append(Paragraph('DOSSIER',h))
    refs=[x for x in p.parcels if isinstance(x,dict) and x.get('section') and x.get('numero')]
    cad=' • '.join(str(x.get('section','')).upper()+' '+re.sub(r'\D','',str(x.get('numero',''))) for x in refs)
    add_table([('Proprietaire',p.owner),('Telephone',p.phone),('Adresse',p.address),('Commune',p.commune),('Cadastre',cad),('Type',p.property_type),('Surface habitable',(p.surface+' m²') if p.surface else ''),('Surface Carrez',(p.surface_carrez+' m²') if p.surface_carrez else '')])
    if p.commune and refs:
        story.append(Paragraph('PLAN CADASTRAL',h))
        try:
            img,_,_=_multi_context_map_png(p.commune,refs,760,500) if len(refs)>1 else (lambda x:(x[0],x[1],x[2]))(_context_map_png(p.commune,refs[0]['section'],refs[0]['numero'],760,500))
            story.append(RLImage(img,width=164*mm,height=108*mm));story.append(Paragraph((p.commune+' - Parcelles : '+cad),cap))
        except Exception: story.append(Paragraph('Plan cadastral temporairement indisponible lors de la generation du document.',cap))
    if p.sections:
        story.append(Paragraph('RELEVE DE VISITE',h))
        for s in p.sections:
            rows=s.get('rows') or []
            rows=[(x[0],x[1]) for x in rows if isinstance(x,(list,tuple)) and len(x)>=2 and str(x[1]).strip()]
            if rows: story.append(Paragraph(str(s.get('title') or 'Informations'),h));add_table(rows)
    if p.photos:
        story.append(Paragraph('PHOTOGRAPHIES DU BIEN',h)); cells=[]
        for item in p.photos[:30]:
            try:
                data=str(item.get('data','')); raw=base64.b64decode(data.split(',',1)[1]); bio=io.BytesIO(raw)
                im=RLImage(bio,width=76*mm,height=52*mm,kind='proportional'); label=Paragraph(str(item.get('name') or 'Photo'),cap)
                cells.append([im,label])
            except Exception: pass
        for i in range(0,len(cells),2):
            row=cells[i:i+2]
            while len(row)<2: row.append(['',''])
            t=Table([[row[0][0],row[1][0]],[row[0][1],row[1][1]]],colWidths=[82*mm,82*mm],hAlign='CENTER')
            t.setStyle(TableStyle([('BOX',(0,0),(-1,-1),.35,colors.HexColor('#CCD5DB')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(0,0),(-1,-1),'CENTER'),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]));story.append(KeepTogether(t));story.append(Spacer(1,2*mm))
    def footer(canvas,docx):
        canvas.saveState();canvas.setFont('Helvetica',7);canvas.setFillColor(colors.HexColor('#777777'));canvas.drawCentredString(A4[0]/2,7*mm,f'ESTIM’IA - Fiche de visite - Page {docx.page}');canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer);out.seek(0)
    safe=re.sub(r'[^A-Za-z0-9_-]+','_',p.owner or p.address or 'bien').strip('_') or 'bien'
    return StreamingResponse(out,media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="Fiche_visite_{safe}.pdf"'})

@app.post("/api/transcribe")
async def transcribe(audio: UploadFile=File(...)):
    content=await audio.read()
    if not content: raise HTTPException(400,"Enregistrement audio vide.")
    ctype=audio.content_type or "audio/webm"
    try:
        result=multipart_request(
            "https://api.openai.com/v1/audio/transcriptions",
            {"model":"gpt-transcribe","language":"fr"},
            [("file",audio.filename or "visite.webm",content,ctype)],
            {"Authorization":f"Bearer {api_key()}"}
        )
        return {"text":result.get("text","")}
    except urllib.error.HTTPError as e:
        detail=e.read().decode(errors="ignore")
        raise HTTPException(e.code,detail[:800])
    except Exception as e: raise HTTPException(500,str(e))

class StructReq(BaseModel):
    notes:str
    address:str=""
    type:str=""
    surface:str=""

@app.post("/api/structure")
def structure(req:StructReq):
    prompt=f"""Tu es ESTIM'IA, assistant d'un agent immobilier français.
Transforme les notes de visite en JSON factuel. N'invente rien.
Conserve les incertitudes avec la mention "à confirmer".
Adresse: {req.address}
Type déclaré: {req.type}
Surface déclarée: {req.surface}
Notes:
{req.notes}

Retourne UNIQUEMENT un objet JSON avec, quand disponible:
type_retenu, surface_m2, pieces, chambres, surfaces_chambres, etage_niveau,
exposition, terrasse_balcon, parking_garage, cave_annexes, etat_general,
travaux_rafraichissement, chauffage, climatisation, menuiseries, volets,
dpe, diagnostics, acces, vue, cuisine, salle_eau_sdb, wc, points_valorisants,
points_a_pondérer, questions_a_confirmer."""
    payload=json.dumps({
        "model":"gpt-5.6-luna",
        "input":prompt,
        "text":{"format":{"type":"json_object"}}
    }).encode()
    request=urllib.request.Request("https://api.openai.com/v1/responses",data=payload,headers={
        "Authorization":f"Bearer {api_key()}","Content-Type":"application/json"
    },method="POST")
    try:
        with urllib.request.urlopen(request,timeout=120) as r: data=json.loads(r.read())
        text=""
        for item in data.get("output",[]):
            for c in item.get("content",[]):
                if c.get("type")=="output_text": text+=c.get("text","")
        return json.loads(text)
    except urllib.error.HTTPError as e:
        raise HTTPException(e.code,e.read().decode(errors="ignore")[:800])
    except Exception as e: raise HTTPException(500,str(e))

@app.post("/api/save")
async def save(address:str=Form(""),type:str=Form(""),surface:str=Form(""),notes:str=Form(""),correction:str=Form(""),
               owner:str=Form(""),owner_phone:str=Form(""),cad_commune:str=Form(""),cad_section:str=Form(""),cad_parcel:str=Form(""),
               photos:list[UploadFile]=File(default=[])):
    did="EST-"+uuid.uuid4().hex[:8].upper()
    folder=DATA/did; folder.mkdir()
    meta={"id":did,"owner":owner,"owner_phone":owner_phone,"address":address,"type":type,"surface":surface,
          "cadastre":{"commune":cad_commune,"section":cad_section,"parcel":cad_parcel},
          "notes":notes,"correction":correction}
    saved=0
    for i,p in enumerate(photos[:60],1):
        raw=await p.read()
        ext=Path(p.filename or "").suffix.lower() or ".jpg"
        (folder/f"{i:02d}{ext}").write_bytes(raw); saved+=1
    (folder/"dossier.json").write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding="utf-8")
    return {"id":did,"photos":saved}
